# create_template.py — 生成支持多条货物的合同模板
# 运行一次即可：python create_template.py

from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = Path("templates") / "采购合同模板.docx"
Path("templates").mkdir(exist_ok=True)

FONT_CN = "宋体"      # 中文字体
FONT_EN = "Times New Roman"  # 西文字体


def _set_run_font(run, size_pt=None, bold=None, cn=FONT_CN, en=FONT_EN):
    """给 run 同时设置中英文字体，避免 Word 自动替换导致方块字"""
    run.font.name = en
    run._r.get_or_add_rPr()
    rFonts = run._r.rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"),    en)
    rFonts.set(qn("w:hAnsi"),    en)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:cs"),       cn)
    if size_pt:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def _set_doc_default_font(doc, cn=FONT_CN, en=FONT_EN):
    """设置文档级别默认字体，作为兜底保障"""
    doc_defaults = doc.styles.element.find(qn("w:docDefaults"))
    if doc_defaults is None:
        return
    rPrDefault = doc_defaults.find(qn("w:rPrDefault"))
    if rPrDefault is None:
        rPrDefault = OxmlElement("w:rPrDefault")
        doc_defaults.append(rPrDefault)
    rPr = rPrDefault.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        rPrDefault.append(rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"),    en)
    rFonts.set(qn("w:hAnsi"),    en)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:cs"),       cn)


def add_run(para, text, size_pt=11, bold=False):
    r = para.add_run(text)
    _set_run_font(r, size_pt=size_pt, bold=bold)
    return r


def bold_para(doc, text, size=12, center=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, text, size_pt=size, bold=True)
    return p


def clause(doc, text, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(24)
    add_run(p, text, size_pt=size)
    return p


def set_cell_text(cell, text, size=11, bold=False):
    """清空单元格并写入带字体的文字"""
    cell.text = ""
    p = cell.paragraphs[0]
    add_run(p, text, size_pt=size, bold=bold)


# ── 文档初始化 ────────────────────────────────────────────────────────
doc = Document()
_set_doc_default_font(doc)

for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(2.5)

# ── 标题 ──────────────────────────────────────────────────────────────
bold_para(doc, "采  购  合  同", size=22, center=True)
doc.add_paragraph()

# ── 编号 & 日期 ───────────────────────────────────────────────────────
p = doc.add_paragraph()
add_run(p, "合同编号：", bold=True)
add_run(p, "{{合同编号}}")
add_run(p, "          签署日期：", bold=True)
add_run(p, "{{签署日期}}")

doc.add_paragraph()

# ── 甲乙方信息 ────────────────────────────────────────────────────────
pt = doc.add_table(rows=4, cols=2)
pt.style = "Table Grid"
rows_data = [
    ("甲方（买方）：{{甲方名称}}", "地址：{{甲方地址}}"),
    ("联系人：{{甲方联系人}}",     "电话：{{甲方电话}}"),
    ("乙方（卖方）：{{乙方名称}}", "地址：{{乙方地址}}"),
    ("联系人：{{乙方联系人}}",     "电话：{{乙方电话}}"),
]
for r, (a, b) in enumerate(rows_data):
    set_cell_text(pt.rows[r].cells[0], a)
    set_cell_text(pt.rows[r].cells[1], b)

doc.add_paragraph()

# ── 一、货物清单（多条，Jinja2 loop）────────────────────────────────
bold_para(doc, "一、货物清单")

COLS  = ["序号", "货物名称及规格", "品牌", "数量", "单位", "单价（元）", "小计（元）"]
COL_W = [1.2,    5.5,              2.5,    1.5,    1.2,    2.5,          2.5]

gt = doc.add_table(rows=5, cols=len(COLS))
gt.style = "Table Grid"

# 表头行
for i, (h, w) in enumerate(zip(COLS, COL_W)):
    c = gt.rows[0].cells[i]
    set_cell_text(c, h, bold=True)
    c.width = Cm(w)

# {%tr for %} 控制行
for_ctrl = ["{%tr for item in items %}"] + [""] * (len(COLS) - 1)
for i, v in enumerate(for_ctrl):
    set_cell_text(gt.rows[1].cells[i], v)

# 数据行
data_row = [
    "{{item.序号}}", "{{item.货物名称}}", "{{item.品牌}}",
    "{{item.数量}}", "{{item.单位}}",     "{{item.单价}}", "{{item.小计}}",
]
for i, v in enumerate(data_row):
    set_cell_text(gt.rows[2].cells[i], v)

# {%tr endfor %} 控制行
endfor_ctrl = ["{%tr endfor %}"] + [""] * (len(COLS) - 1)
for i, v in enumerate(endfor_ctrl):
    set_cell_text(gt.rows[3].cells[i], v)

# 合计行
total_row = ["", "合    计", "", "", "", "", "{{合同金额数字}} 元"]
for i, v in enumerate(total_row):
    set_cell_text(gt.rows[4].cells[i], v, bold=(i in (1, 6)))

doc.add_paragraph()

# ── 附：设备组成明细（条件显示，breakdown_items 非空时渲染）────────────
# docxtpl 条件块：{%p if breakdown_items %} ... {%p endif %}
p_if = doc.add_paragraph()
p_if.paragraph_format.space_before = Pt(0)
p_if.paragraph_format.space_after  = Pt(0)
add_run(p_if, "{%p if breakdown_items %}", size_pt=1)  # 不可见占位行

bold_para(doc, "附：设备组成明细")

BD_COLS  = ["序号", "组件名称", "数量", "单价（元）", "小计（元）"]
BD_COL_W = [1.2,    8.5,        1.5,    2.5,           2.5]

bdt = doc.add_table(rows=4, cols=len(BD_COLS))
bdt.style = "Table Grid"

for i, (h, w) in enumerate(zip(BD_COLS, BD_COL_W)):
    c = bdt.rows[0].cells[i]
    set_cell_text(c, h, bold=True)
    c.width = Cm(w)

bd_for = ["{%tr for bd in breakdown_items %}"] + [""] * (len(BD_COLS) - 1)
for i, v in enumerate(bd_for):
    set_cell_text(bdt.rows[1].cells[i], v)

bd_data = ["{{bd.序号}}", "{{bd.组件名称}}", "{{bd.数量}}", "{{bd.单价}}", "{{bd.小计}}"]
for i, v in enumerate(bd_data):
    set_cell_text(bdt.rows[2].cells[i], v)

bd_endfor = ["{%tr endfor %}"] + [""] * (len(BD_COLS) - 1)
for i, v in enumerate(bd_endfor):
    set_cell_text(bdt.rows[3].cells[i], v)

doc.add_paragraph()

p_endif = doc.add_paragraph()
p_endif.paragraph_format.space_before = Pt(0)
p_endif.paragraph_format.space_after  = Pt(0)
add_run(p_endif, "{%p endif %}", size_pt=1)

# ── 二、合同金额 ──────────────────────────────────────────────────────
bold_para(doc, "二、合同金额")
clause(doc, "合同总价款人民币（大写）{{合同金额大写}}，小写 ¥{{合同金额数字}} 元，"
            "含税率 {{税率}}%，税额 {{税额}} 元。")

# ── 三、交货条款 ──────────────────────────────────────────────────────
bold_para(doc, "三、交货条款")
clause(doc, "1. 交货地点：{{交货地点}}。")
clause(doc, "2. 交货期限：合同签订后 {{交货期限}} 天内完成交付；运费由{{运费承担方}}承担。")
clause(doc, "3. 质量标准：{{质量标准}}；质保期 {{质保期}} 个月。")

# ── 四、付款方式 ──────────────────────────────────────────────────────
bold_para(doc, "四、付款方式")
clause(doc, "1. 合同签订后 {{预付款天数}} 个工作日内，甲方支付预付款（{{预付款比例}}%）"
            "共 ¥{{预付款金额}} 元。")
clause(doc, "2. 出厂前验收合格后，甲方支付出厂验收款（{{出厂验收比例}}%）共 ¥{{出厂验收金额}} 元。")
clause(doc, "3. 到厂验收合格后 {{到货付款天数}} 个工作日内，甲方支付尾款（{{到厂验收比例}}%）"
            "共 ¥{{到厂验收金额}} 元。")
clause(doc, "4. 收款户名：{{收款户名}}；开户行：{{开户银行}}；账号：{{银行账号}}。")

# ── 五、验收与违约 ────────────────────────────────────────────────────
bold_para(doc, "五、验收与违约")
clause(doc, "1. 甲方收货后 {{验收期限}} 天内完成验收；质量问题乙方 {{整改期限}} 天内整改。")
clause(doc, "2. 逾期交货：每日按合同金额 {{逾期违约金比例}}‰ 支付违约金。")
clause(doc, "3. 逾期付款：每日按逾期金额 {{逾期付款利率}}‰ 支付利息。")

# ── 六、其他约定 ──────────────────────────────────────────────────────
bold_para(doc, "六、其他约定")
clause(doc, "{{其他约定}}")
clause(doc, "争议提交{{争议解决地}}仲裁委员会仲裁。本合同一式两份，签字盖章后生效。")

doc.add_paragraph()
doc.add_paragraph()

# ── 签字页 ────────────────────────────────────────────────────────────
st = doc.add_table(rows=3, cols=2)
st.style = "Table Grid"
sig_rows = [
    ("甲方（盖章）：", "乙方（盖章）："),
    ("授权代表（签字）：", "授权代表（签字）："),
    ("日期：", "日期："),
]
for r, (a, b) in enumerate(sig_rows):
    set_cell_text(st.rows[r].cells[0], a)
    set_cell_text(st.rows[r].cells[1], b)

doc.save(OUTPUT)
print(f"✓ 新模板已生成：{OUTPUT}（字体：{FONT_CN} / {FONT_EN}）")
