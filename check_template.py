# check_template.py — 核查模板占位符与代码字段的对齐情况
# 运行：python check_template.py

import re
from docx import Document
from mock_data import auto_generate, MOCK_VOICE_INPUT

def extract_placeholders(docx_path):
    doc = Document(docx_path)
    found = set()
    for para in doc.paragraphs:
        for m in re.findall(r'\{\{([^}]+)\}\}', para.text):
            found.add(m.strip())
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for m in re.findall(r'\{\{([^}]+)\}\}', para.text):
                        found.add(m.strip())
    return found

template_fields = extract_placeholders("templates/采购合同模板.docx")
code_fields     = set(auto_generate(MOCK_VOICE_INPUT).keys())

missing_in_code  = template_fields - code_fields   # 模板有，代码没给值
extra_in_code    = code_fields - template_fields    # 代码有，模板没用到

print(f"模板占位符共 {len(template_fields)} 个：")
for f in sorted(template_fields):
    print(f"  {{{{ {f} }}}}")

print(f"\n⚠️  模板有、代码未提供（会留空白）：{missing_in_code or '无'}")
print(f"ℹ️  代码有、模板未使用（多余字段）：{extra_in_code or '无'}")
